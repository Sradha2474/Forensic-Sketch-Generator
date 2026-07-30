"""
Build STAAR assessment PowerPoint for Forensic Sketch Generator.
Follows MP_ppt_format.pptx theme (Times New Roman, accent blues).
"""

from __future__ import annotations

import copy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt, Emu
from pptx.oxml import parse_xml
from lxml import etree

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
TEMPLATE = ROOT / "MP_ppt_format.pptx"
LOGO = ROOT / "image1.png"
OUT = ROOT / "Forensic_Sketch_Generator_STAAR_Presentation.pptx"

# Template theme colors
ACCENT1 = RGBColor(0x5B, 0x9B, 0xD5)  # primary blue
ACCENT5 = RGBColor(0x44, 0x72, 0xC4)  # deeper blue
DK2 = RGBColor(0x44, 0x54, 0x6A)      # slate
LT2 = RGBColor(0xE7, 0xE6, 0xE6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
ORANGE = RGBColor(0xED, 0x7D, 0x31)
GREEN = RGBColor(0x70, 0xAD, 0x47)
GOLD = RGBColor(0xFF, 0xC0, 0x00)
LIGHT_BG = RGBColor(0xF2, 0xF7, 0xFC)
SOFT_BLUE = RGBColor(0xDE, 0xEB, 0xF7)
NAVY = RGBColor(0x1F, 0x4E, 0x79)

FONT = "Times New Roman"


def _delete_all_slides(prs: Presentation) -> None:
    """Remove every slide while keeping slide master / theme."""
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        rId = sldId.get(qn("r:id"))
        prs.part.drop_rel(rId)
        sldIdLst.remove(sldId)


def _set_run(run, text, size=18, bold=False, color=BLACK, font=FONT):
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    # Force East Asian / Latin font
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", font)


def _add_text_box(slide, left, top, width, height, text, size=18, bold=False,
                  color=BLACK, align=PP_ALIGN.LEFT, font=FONT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    _set_run(run, text, size=size, bold=bold, color=color, font=font)
    return box


def _add_bullets(slide, left, top, width, height, items, size=16, color=BLACK, spacing=8):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.level = 0
        p.space_after = Pt(spacing)
        run = p.add_run()
        _set_run(run, f"•  {item}", size=size, color=color)
    return box


def _header_bar(slide, title: str, icon_name: str | None = None):
    """Blue header matching Literature Study slide style."""
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.05)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT5
    bar.line.fill.background()

    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.05), Inches(13.333), Inches(0.06)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = ACCENT1
    accent.line.fill.background()

    _add_text_box(
        slide, Inches(0.55), Inches(0.22), Inches(11.5), Inches(0.7),
        title, size=28, bold=True, color=WHITE
    )
    if icon_name:
        icon = ASSETS / f"{icon_name}.png"
        if icon.exists():
            slide.shapes.add_picture(str(icon), Inches(12.15), Inches(0.15), Inches(0.75), Inches(0.75))


def _footer(slide, page: int, total: int = 20):
    _add_text_box(
        slide, Inches(0.4), Inches(7.05), Inches(10), Inches(0.35),
        "STAAR Assessment  |  Forensic Sketch Generator  |  C.V. Raman Global University",
        size=10, color=DK2
    )
    _add_text_box(
        slide, Inches(12.2), Inches(7.05), Inches(0.9), Inches(0.35),
        str(page), size=12, bold=True, color=ACCENT5, align=PP_ALIGN.RIGHT
    )


def _card(slide, left, top, width, height, fill=SOFT_BLUE, line=ACCENT1):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line
    shp.line.width = Pt(1.25)
    try:
        shp.adjustments[0] = 0.1
    except Exception:
        pass
    return shp


def _process_box(slide, left, top, width, height, title, subtitle, fill):
    shp = _card(slide, left, top, width, height, fill=fill, line=fill)
    _add_text_box(slide, left + Inches(0.08), top + Inches(0.18), width - Inches(0.1), Inches(0.4),
                  title, size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    _add_text_box(slide, left + Inches(0.08), top + Inches(0.55), width - Inches(0.1), Inches(0.4),
                  subtitle, size=11, color=WHITE, align=PP_ALIGN.CENTER)
    return shp


def _arrow(slide, left, top):
    shp = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, Inches(0.28), Inches(0.22))
    shp.fill.solid()
    shp.fill.fore_color.rgb = DK2
    shp.line.fill.background()
    return shp


def _set_cell(cell, text, size=11, bold=False, color=BLACK, fill=None, align=PP_ALIGN.LEFT):
    cell.text = ""
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    _set_run(run, text, size=size, bold=bold, color=color)
    if fill is not None:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        # remove existing solid fill
        for child in list(tcPr):
            if "solidFill" in child.tag:
                tcPr.remove(child)
        solid = etree.SubElement(tcPr, qn("a:solidFill"))
        srgb = etree.SubElement(solid, qn("a:srgbClr"))
        srgb.set("val", f"{fill[0]:02X}{fill[1]:02X}{fill[2]:02X}")


def build():
    prs = Presentation(str(TEMPLATE))
    _delete_all_slides(prs)
    # Keep slide size from template
    blank = prs.slide_layouts[6]  # Blank
    title_layout = prs.slide_layouts[0]
    content_layout = prs.slide_layouts[1]

    # ========== 1. TITLE ==========
    slide = prs.slides.add_slide(blank)
    # Top banner
    top = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.15), Inches(0.05), Inches(13.0), Inches(0.85))
    top.fill.solid()
    top.fill.fore_color.rgb = ACCENT1
    top.line.fill.background()
    _add_text_box(slide, Inches(0.3), Inches(0.18), Inches(12.5), Inches(0.6),
                  "Major Project Presentation", size=32, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # Title band
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.35), Inches(1.05), Inches(12.6), Inches(1.35))
    band.fill.solid()
    band.fill.fore_color.rgb = ACCENT5
    band.line.fill.background()
    _add_text_box(slide, Inches(0.5), Inches(1.15), Inches(12.3), Inches(0.4),
                  "On", size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    _add_text_box(slide, Inches(0.5), Inches(1.5), Inches(12.3), Inches(0.7),
                  "Forensic Sketch Generator: An AI-Powered Approach\nUsing NLP and Conditional GAN",
                  size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    if LOGO.exists():
        slide.shapes.add_picture(str(LOGO), Inches(5.7), Inches(2.6), Inches(1.9), Inches(1.7))

    _add_text_box(slide, Inches(0.8), Inches(4.45), Inches(11.7), Inches(1.5),
                  "1. Tusharkanta Behera    Regd. No.: 2301020601\n"
                  "2. Sradha Ram                Regd. No.: 2301020808\n"
                  "3. Sandeep Kumar Swain  Regd. No.: 2301020438",
                  size=15, color=DK2, align=PP_ALIGN.CENTER)

    _add_text_box(slide, Inches(0.8), Inches(5.95), Inches(11.7), Inches(0.9),
                  "Under the supervision of Dr. Sukant Kishoro Bisoy\n"
                  "Department of CSE  |  C.V. Raman Global University, Bhubaneswar\n"
                  "Science and Technology Advancement — Assessment Review (STAAR)",
                  size=13, color=NAVY, align=PP_ALIGN.CENTER)

    # ========== 2. PROBLEM STATEMENT ==========
    slide = prs.slides.add_slide(blank)
    _header_bar(slide, "Problem Statement", "problem")
    hero = ASSETS / "img_witness.jpg"
    if hero.exists():
        slide.shapes.add_picture(str(hero), Inches(8.6), Inches(1.4), Inches(4.3), Inches(5.2))
        # overlay label
        ov = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.6), Inches(6.1), Inches(4.3), Inches(0.5))
        ov.fill.solid()
        ov.fill.fore_color.rgb = ACCENT5
        ov.line.fill.background()
        _add_text_box(slide, Inches(8.7), Inches(6.15), Inches(4.1), Inches(0.4),
                      "Traditional composite creation", size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    _card(slide, Inches(0.4), Inches(1.4), Inches(7.9), Inches(5.3), fill=LIGHT_BG, line=ACCENT1)
    _add_bullets(slide, Inches(0.7), Inches(1.6), Inches(7.4), Inches(4.9), [
        "Facial composites begin with a Cognitive Interview — hours of recall, catalog picking, and revision.",
        "Trained forensic sketch artists are scarce, costly, and introduce subjective bias.",
        "Hand and software composites (Identikit, E-FIT, FACES) still demand expert operators.",
        "Verbal memory often fails to map cleanly onto fixed feature libraries.",
        "Early investigation windows need usable likenesses in minutes, not multi-hour sessions.",
        "Gap: no artist-aligned AI that interviews → structures → sketches → refines with the witness.",
    ], size=15, color=DK2, spacing=10)
    _footer(slide, 2)

    # ========== 3. MOTIVATION ==========
    slide = prs.slides.add_slide(blank)
    _header_bar(slide, "Motivation", "motivation")
    cards = [
        ("Speed", "Cut composite turnaround from hours to minutes for time-critical leads."),
        ("Access", "Assist investigators where certified artists are unavailable."),
        ("Fidelity", "Keep witness memory as the source of truth — not unconstrained prompts."),
        ("Practice Fit", "Digitize CI / H-CI workflow instead of generic text-to-image."),
        ("Control", "Edit one facial trait at a time without discarding identity."),
        ("Research", "Enable measurable quality via FID, CLIP, and ArcFace metrics."),
    ]
    colors = [ACCENT5, ACCENT1, ORANGE, GREEN, GOLD, DK2]
    for i, ((t, b), c) in enumerate(zip(cards, colors)):
        r, col = divmod(i, 3)
        left = Inches(0.45 + col * 4.2)
        top = Inches(1.4 + r * 2.6)
        shp = _card(slide, left, top, Inches(3.95), Inches(2.3), fill=SOFT_BLUE, line=c)
        badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, left + Inches(1.45), top + Inches(0.2), Inches(0.9), Inches(0.9))
        badge.fill.solid()
        badge.fill.fore_color.rgb = c
        badge.line.fill.background()
        _add_text_box(slide, left + Inches(1.45), top + Inches(0.4), Inches(0.9), Inches(0.5),
                      str(i + 1), size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        _add_text_box(slide, left + Inches(0.2), top + Inches(1.2), Inches(3.55), Inches(0.4),
                      t, size=16, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        _add_text_box(slide, left + Inches(0.2), top + Inches(1.6), Inches(3.55), Inches(0.55),
                      b, size=12, color=DK2, align=PP_ALIGN.CENTER)
    _footer(slide, 3)

    # ========== 4. EXISTING METHODS ==========
    slide = prs.slides.add_slide(blank)
    _header_bar(slide, "Existing Methods", "existing")
    rows_data = [
        ("Method", "How it works", "Limitation"),
        ("Hand sketch + CI", "Artist interviews & draws", "Scarce talent; hours per session"),
        ("Identikit / E-FIT / FACES", "Catalog feature assembly", "Needs expert; library gaps"),
        ("EvoFIT / EFIT-V", "Holistic evolutionary faces", "Still operator-guided software"),
        ("Text-to-Face GANs", "Caption → photo face", "Photo output; weak interview loop"),
        ("Sketch → Photo nets", "Translate existing sketches", "Assumes sketch already exists"),
    ]
    table = slide.shapes.add_table(len(rows_data), 3, Inches(0.5), Inches(1.4), Inches(12.3), Inches(5.2)).table
    table.columns[0].width = Inches(3.2)
    table.columns[1].width = Inches(4.6)
    table.columns[2].width = Inches(4.5)
    for r, row in enumerate(rows_data):
        for c, val in enumerate(row):
            if r == 0:
                _set_cell(table.cell(r, c), val, size=13, bold=True, color=WHITE,
                          fill=(0x44, 0x72, 0xC4), align=PP_ALIGN.CENTER)
            else:
                fill = (0xF2, 0xF7, 0xFC) if r % 2 else (0xFF, 0xFF, 0xFF)
                _set_cell(table.cell(r, c), val, size=12, color=DK2, fill=fill)
    _footer(slide, 4)

    # ========== 5. LITERATURE SURVEY ==========
    slide = prs.slides.add_slide(blank)
    _header_bar(slide, "Literature Survey", "literature")
    lit = [
        ("Sl.", "Venue / Year", "Work", "Method", "Finding", "Gap"),
        ("1", "WJARR 2026", "Soppari et al.\n(Base paper)", "SDXL → InsightFace\n→ FAISS retrieval", "Text→sketch→rank\nreduces artist need", "Weak iterative\nrefinement"),
        ("2", "CVPR / arXiv", "TediGAN; ST²FG", "Text→StyleGAN;\nBERT + ACM edits", "High-res control;\nuser refinement", "Photo-like; GAN\ninstability"),
        ("3", "Elsevier / IJCB", "DCGAN STF;\nCLIP4Sketch", "Sketch↔photo &\nCLIP matching", "Strong once sketch\nexists", "No text-to-\ncomposite start"),
        ("4", "IEEE SPL 2017", "Galea & Farrugia", "VGG-Face +\n3DMM augment", "94.61% Rank-1\nsketch–photo", "Needs sketch;\nno generation"),
    ]
    table = slide.shapes.add_table(len(lit), 6, Inches(0.25), Inches(1.35), Inches(12.8), Inches(5.3)).table
    widths = [0.6, 1.8, 2.2, 2.8, 2.7, 2.7]
    for i, w in enumerate(widths):
        table.columns[i].width = Inches(w)
    for r, row in enumerate(lit):
        for c, val in enumerate(row):
            if r == 0:
                _set_cell(table.cell(r, c), val, size=11, bold=True, color=WHITE,
                          fill=(0x44, 0x72, 0xC4), align=PP_ALIGN.CENTER)
            else:
                fill = (0xDE, 0xEB, 0xF7) if r % 2 else (0xFF, 0xFF, 0xFF)
                _set_cell(table.cell(r, c), val, size=10, color=DK2, fill=fill,
                          align=PP_ALIGN.CENTER if c == 0 else PP_ALIGN.LEFT)
    _footer(slide, 5)

    # ========== 6. RESEARCH GAP ==========
    slide = prs.slides.add_slide(blank)
    _header_bar(slide, "Research Gap", "gap")
    gap_rows = [
        ("Dimension", "Current State", "Our Vision"),
        ("Starting point", "Needs hand / software composite first", "AI-led forensic interview → first draft"),
        ("Workflow", "One-shot prompt generation", "Interview → profile → construct → refine"),
        ("Feature control", "Opaque prompt / full regeneration", "Living profile; edit one trait at a time"),
        ("Output form", "Photo faces or sketch-to-photo only", "Explicit forensic pencil-style sketch"),
        ("Product role", "Artist replacement framing", "Investigator assist; witness authority"),
    ]
    table = slide.shapes.add_table(len(gap_rows), 3, Inches(0.4), Inches(1.35), Inches(12.5), Inches(4.5)).table
    table.columns[0].width = Inches(2.4)
    table.columns[1].width = Inches(5.05)
    table.columns[2].width = Inches(5.05)
    for r, row in enumerate(gap_rows):
        for c, val in enumerate(row):
            if r == 0:
                _set_cell(table.cell(r, c), val, size=13, bold=True, color=WHITE,
                          fill=(0x44, 0x72, 0xC4), align=PP_ALIGN.CENTER)
            else:
                if c == 0:
                    _set_cell(table.cell(r, c), val, size=12, bold=True, color=WHITE, fill=(0x5B, 0x9B, 0xD5))
                elif c == 1:
                    _set_cell(table.cell(r, c), val, size=12, color=DK2, fill=(0xFD, 0xE9, 0xD9))
                else:
                    _set_cell(table.cell(r, c), val, size=12, color=DK2, fill=(0xE2, 0xEF, 0xDA))
    _card(slide, Inches(0.4), Inches(6.0), Inches(12.5), Inches(0.85), fill=SOFT_BLUE, line=ACCENT5)
    _add_text_box(slide, Inches(0.6), Inches(6.15), Inches(12.1), Inches(0.55),
                  "Contribution: unify artist-style interviewing, controllable face construction, forensic sketch rendering, "
                  "witness-in-the-loop refinement, and research-time biometric evaluation.",
                  size=13, bold=True, color=NAVY)
    _footer(slide, 6)

    # ========== 7. PROPOSED SOLUTION ==========
    slide = prs.slides.add_slide(blank)
    _header_bar(slide, "Proposed Solution", "solution")
    _add_text_box(slide, Inches(0.5), Inches(1.3), Inches(12.3), Inches(0.45),
                  "Digitize composite practice: interview → structure features → compose face → render sketch → refine",
                  size=15, bold=True, color=NAVY, align=PP_ALIGN.CENTER)

    modules = [
        ("A", "Interview Agent", "CI / H-CI protocol\nstructured dialogue", ACCENT5),
        ("B", "Feature Profile", "Living attribute\nschema (13+ traits)", ACCENT1),
        ("C", "Face Composer", "Controllable gen\nStyleGAN2 / SD", ORANGE),
        ("D", "Sketch Renderer", "Pencil forensic\ncomposite styling", GREEN),
        ("E", "Refinement", "Selective edits\nwitness-approved", GOLD),
    ]
    for i, (letter, title, desc, col) in enumerate(modules):
        left = Inches(0.35 + i * 2.55)
        _card(slide, left, Inches(2.0), Inches(2.4), Inches(3.4), fill=LIGHT_BG, line=col)
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, left + Inches(0.75), Inches(2.2), Inches(0.9), Inches(0.9))
        circ.fill.solid()
        circ.fill.fore_color.rgb = col
        circ.line.fill.background()
        _add_text_box(slide, left + Inches(0.75), Inches(2.4), Inches(0.9), Inches(0.5),
                      letter, size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        _add_text_box(slide, left + Inches(0.1), Inches(3.25), Inches(2.2), Inches(0.45),
                      title, size=14, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        _add_text_box(slide, left + Inches(0.1), Inches(3.8), Inches(2.2), Inches(1.2),
                      desc, size=12, color=DK2, align=PP_ALIGN.CENTER)
        if i < 4:
            _arrow(slide, left + Inches(2.35), Inches(3.4))

    _add_text_box(slide, Inches(0.5), Inches(5.7), Inches(12.3), Inches(1.0),
                  "Assistive investigative tool — does not replace witness authority or legal identification.\n"
                  "Optional validation layer: FID · attribute consistency · ArcFace / FaceNet similarity.",
                  size=13, color=DK2, align=PP_ALIGN.CENTER)
    _footer(slide, 7)

    # ========== 8. END-TO-END WORKFLOW ==========
    slide = prs.slides.add_slide(blank)
    _header_bar(slide, "End-to-End Workflow Diagram", "workflow")
    wf = ASSETS / "workflow_diagram.png"
    if wf.exists():
        slide.shapes.add_picture(str(wf), Inches(0.35), Inches(1.35), Inches(12.6), Inches(3.3))

    steps = [
        "Session start with witness",
        "Modules A–B build feature profile",
        "Modules C–D synthesize & sketch",
        "Module E refines until approval",
        "Export approved composite",
        "Optional research validation",
    ]
    for i, s in enumerate(steps):
        r, c = divmod(i, 3)
        left = Inches(0.45 + c * 4.2)
        top = Inches(4.85 + r * 0.95)
        _card(slide, left, top, Inches(4.0), Inches(0.8), fill=SOFT_BLUE, line=ACCENT1)
        _add_text_box(slide, left + Inches(0.15), top + Inches(0.2), Inches(3.7), Inches(0.45),
                      f"{i+1}.  {s}", size=13, bold=True, color=NAVY)
    _footer(slide, 8)

    # ========== 9. SYSTEM ARCHITECTURE ==========
    slide = prs.slides.add_slide(blank)
    _header_bar(slide, "System Architecture Diagram", "architecture")
    arch = ASSETS / "architecture_diagram.png"
    if arch.exists():
        slide.shapes.add_picture(str(arch), Inches(0.35), Inches(1.25), Inches(12.6), Inches(5.5))
    _footer(slide, 9)

    # ========== 10. TECHNOLOGY STACK ==========
    slide = prs.slides.add_slide(blank)
    _header_bar(slide, "Technology Stack", "tech")
    tech_img = ASSETS / "img_tech.jpg"
    if tech_img.exists():
        slide.shapes.add_picture(str(tech_img), Inches(9.0), Inches(1.4), Inches(3.9), Inches(5.2))
        ov = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.0), Inches(6.1), Inches(3.9), Inches(0.5))
        ov.fill.solid()
        ov.fill.fore_color.rgb = ACCENT5
        ov.line.fill.background()
        _add_text_box(slide, Inches(9.1), Inches(6.15), Inches(3.7), Inches(0.4),
                      "PyTorch · Diffusers · CV", size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    stack = [
        ("Client UI", "Next.js (planned)  ·  Streamlit MVP (live)"),
        ("API / Services", "Next.js Route Handlers / FastAPI inference"),
        ("ML Runtime", "Python 3.10+  ·  PyTorch 2.0+  ·  Diffusers"),
        ("Language / NLP", "Protocol-constrained LLM  ·  Transformers"),
        ("Vision Models", "StyleGAN2-ADA  ·  SD v1.5  ·  OpenCV / PIL"),
        ("Validation", "ArcFace / FaceNet  ·  FID / CLIP scores"),
        ("Data & Ops", "CelebA · FFHQ · CUFS  ·  Docker + GPU worker"),
    ]
    for i, (k, v) in enumerate(stack):
        top = Inches(1.35 + i * 0.72)
        _card(slide, Inches(0.4), top, Inches(2.6), Inches(0.62), fill=ACCENT5, line=ACCENT5)
        _add_text_box(slide, Inches(0.5), top + Inches(0.12), Inches(2.4), Inches(0.4),
                      k, size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        _card(slide, Inches(3.15), top, Inches(5.6), Inches(0.62), fill=LIGHT_BG, line=ACCENT1)
        _add_text_box(slide, Inches(3.3), top + Inches(0.12), Inches(5.3), Inches(0.4),
                      v, size=13, color=DK2)
    _footer(slide, 10)

    # ========== 11. AI PIPELINE ==========
    slide = prs.slides.add_slide(blank)
    _header_bar(slide, "AI Pipeline", "ai")
    pipe = ASSETS / "ai_pipeline.png"
    if pipe.exists():
        slide.shapes.add_picture(str(pipe), Inches(0.3), Inches(1.25), Inches(12.7), Inches(4.0))

    _card(slide, Inches(0.4), Inches(5.45), Inches(12.5), Inches(1.35), fill=SOFT_BLUE, line=ACCENT5)
    _add_bullets(slide, Inches(0.65), Inches(5.55), Inches(12.0), Inches(1.2), [
        "MVP path: Witness attributes → JSON profile → prompt builder → Stable Diffusion face → OpenCV pencil sketch.",
        "Research path: LLM interview agent → attribute–latent map → StyleGAN2-ADA → Pix2Pix/CycleGAN sketch → guided refinement.",
        "Unknown attributes stay unknown — no forced hallucination of unstated features.",
    ], size=13, color=DK2, spacing=4)
    _footer(slide, 11)

    # ========== 12. DATASETS ==========
    slide = prs.slides.add_slide(blank)
    _header_bar(slide, "Datasets", "datasets")
    ds = [
        ("Dataset", "Role in Pipeline", "Why it matters"),
        ("CelebA", "Attribute vocabulary & conditional supervision", "Maps language traits to facial labels"),
        ("FFHQ", "High-fidelity face generation prior", "Photoreal identity / quality backbone"),
        ("CUFS / CUFSF", "Photo–sketch translation & evaluation", "Train / test forensic pencil domain"),
        ("Witness profiles (MVP)", "Structured interview JSON from live sessions", "End-to-end prototype validation"),
    ]
    table = slide.shapes.add_table(len(ds), 3, Inches(0.5), Inches(1.4), Inches(12.3), Inches(4.2)).table
    table.columns[0].width = Inches(3.0)
    table.columns[1].width = Inches(5.0)
    table.columns[2].width = Inches(4.3)
    for r, row in enumerate(ds):
        for c, val in enumerate(row):
            if r == 0:
                _set_cell(table.cell(r, c), val, size=13, bold=True, color=WHITE,
                          fill=(0x44, 0x72, 0xC4), align=PP_ALIGN.CENTER)
            else:
                fill = (0xDE, 0xEB, 0xF7) if r % 2 else (0xFF, 0xFF, 0xFF)
                _set_cell(table.cell(r, c), val, size=13, color=DK2, fill=fill)
    _add_text_box(slide, Inches(0.5), Inches(5.9), Inches(12.3), Inches(0.8),
                  "Fairness note: public face corpora may under-represent demographics and rare marks (scars, tattoos) — "
                  "addressed in future fairness analysis.",
                  size=13, color=NAVY, align=PP_ALIGN.CENTER)
    _footer(slide, 12)

    # ========== 13. CURRENT PROGRESS ==========
    slide = prs.slides.add_slide(blank)
    _header_bar(slide, "Current Progress (≈ 50% MVP)", "progress")
    done = [
        ("✓", "Streamlit investigator UI with structured witness interview form", GREEN),
        ("✓", "WitnessProfile JSON schema + save/load to outputs/", GREEN),
        ("✓", "Attribute → Stable Diffusion prompt builder", GREEN),
        ("✓", "SD v1.5 face generation (CPU/GPU) with model cache", GREEN),
        ("✓", "OpenCV edge-preserving forensic pencil sketch + download", GREEN),
        ("✓", "Extension stubs: interview agent, selective edit, eval, DB search", GOLD),
    ]
    for i, (mark, text, col) in enumerate(done):
        top = Inches(1.35 + i * 0.8)
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.55), top + Inches(0.05), Inches(0.55), Inches(0.55))
        circ.fill.solid()
        circ.fill.fore_color.rgb = col
        circ.line.fill.background()
        _add_text_box(slide, Inches(0.55), top + Inches(0.12), Inches(0.55), Inches(0.4),
                      mark, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        _card(slide, Inches(1.3), top, Inches(11.4), Inches(0.65), fill=LIGHT_BG, line=col)
        _add_text_box(slide, Inches(1.5), top + Inches(0.15), Inches(11.0), Inches(0.4),
                      text, size=14, color=DK2)
    _footer(slide, 13)

    # ========== 14. FUTURE WORK ==========
    slide = prs.slides.add_slide(blank)
    _header_bar(slide, "Future Work", "future")
    futures = [
        ("Conversational Agent", "Replace form with protocol-constrained LLM interview (Module A)."),
        ("StyleGAN2-ADA Control", "Attribute-level latent edits without identity drift."),
        ("Sketch Translation", "Pix2Pix / CycleGAN on CUFS for stronger pencil domain."),
        ("Selective Inpainting", "Region edits for scars, spacing, jaw — ControlNet hooks."),
        ("Evaluation Suite", "FID, CLIP attribute score, ArcFace identity metrics."),
        ("Practitioner Studies", "Controlled trials with forensic artists & witness-role users."),
    ]
    for i, (t, d_) in enumerate(futures):
        r, c = divmod(i, 3)
        left = Inches(0.4 + c * 4.25)
        top = Inches(1.4 + r * 2.55)
        _card(slide, left, top, Inches(4.05), Inches(2.3), fill=SOFT_BLUE, line=ORANGE)
        num = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.2), top + Inches(0.25),
                                     Inches(0.55), Inches(0.55))
        num.fill.solid()
        num.fill.fore_color.rgb = ORANGE
        num.line.fill.background()
        _add_text_box(slide, left + Inches(0.2), top + Inches(0.32), Inches(0.55), Inches(0.4),
                      str(i + 1), size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        _add_text_box(slide, left + Inches(0.9), top + Inches(0.3), Inches(2.95), Inches(0.5),
                      t, size=14, bold=True, color=NAVY)
        _add_text_box(slide, left + Inches(0.25), top + Inches(1.05), Inches(3.55), Inches(1.0),
                      d_, size=13, color=DK2)
    _footer(slide, 14)

    # ========== 15. EXPECTED IMPACT ==========
    slide = prs.slides.add_slide(blank)
    _header_bar(slide, "Expected Impact", "impact")
    impacts = [
        ("Minutes, not hours", "Usable forensic-style composites for early investigation windows."),
        ("Procedural consistency", "Protocol-constrained interviewing reduces operator drift."),
        ("Wider access", "Assist units lacking on-call composite artists."),
        ("Measurable quality", "Research metrics enable objective iteration."),
        ("Ethical framing", "Witness-in-the-loop; assistive — not courtroom evidence."),
        ("Open research path", "Modular design allows model swap without UI rewrite."),
    ]
    for i, (t, b) in enumerate(impacts):
        r, c = divmod(i, 3)
        left = Inches(0.4 + c * 4.25)
        top = Inches(1.4 + r * 2.55)
        _card(slide, left, top, Inches(4.05), Inches(2.3), fill=LIGHT_BG, line=GREEN)
        _add_text_box(slide, left + Inches(0.2), top + Inches(0.35), Inches(3.65), Inches(0.55),
                      t, size=16, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        _add_text_box(slide, left + Inches(0.25), top + Inches(1.1), Inches(3.55), Inches(0.9),
                      b, size=13, color=DK2, align=PP_ALIGN.CENTER)
    _footer(slide, 15)

    # ========== 16. DEMO / PROTOTYPE ==========
    slide = prs.slides.add_slide(blank)
    _header_bar(slide, "Demo / Prototype", "demo")
    demo = ASSETS / "demo_ui_mockup.png"
    if demo.exists():
        slide.shapes.add_picture(str(demo), Inches(0.35), Inches(1.25), Inches(8.5), Inches(5.0))

    _card(slide, Inches(9.05), Inches(1.25), Inches(3.9), Inches(5.0), fill=SOFT_BLUE, line=ACCENT5)
    _add_text_box(slide, Inches(9.2), Inches(1.45), Inches(3.6), Inches(0.4),
                  "Live MVP Highlights", size=16, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    _add_bullets(slide, Inches(9.25), Inches(2.0), Inches(3.5), Inches(3.5), [
        "Witness interview board",
        "JSON profile export",
        "SD v1.5 generation",
        "Pencil sketch styling",
        "PNG + JSON download",
        "CPU & GPU support",
    ], size=13, color=DK2, spacing=8)
    _add_text_box(slide, Inches(9.2), Inches(5.4), Inches(3.6), Inches(0.6),
                  "github.com/Sradha2474/\nForensic-Sketch-Generator",
                  size=11, bold=True, color=ACCENT5, align=PP_ALIGN.CENTER)
    _footer(slide, 16)

    # ========== 17. TIMELINE ==========
    slide = prs.slides.add_slide(blank)
    _header_bar(slide, "Timeline", "timeline")
    # horizontal timeline
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(3.5), Inches(11.7), Inches(0.08))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT5
    line.line.fill.background()

    phases = [
        ("Phase 1", "Lit. survey &\nproblem framing", "Done", GREEN),
        ("Phase 2", "MVP pipeline\nStreamlit + SD", "Done", GREEN),
        ("Phase 3", "Sketch quality &\nevaluation hooks", "In progress", GOLD),
        ("Phase 4", "Interview agent &\nselective edit", "Next", ORANGE),
        ("Phase 5", "StyleGAN + CUFS\n+ user study", "Planned", ACCENT1),
    ]
    for i, (ph, desc, status, col) in enumerate(phases):
        left = Inches(0.7 + i * 2.45)
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, left + Inches(0.75), Inches(3.28), Inches(0.5), Inches(0.5))
        circ.fill.solid()
        circ.fill.fore_color.rgb = col
        circ.line.fill.background()
        _card(slide, left, Inches(1.5), Inches(2.2), Inches(1.55), fill=LIGHT_BG, line=col)
        _add_text_box(slide, left + Inches(0.1), Inches(1.6), Inches(2.0), Inches(0.35),
                      ph, size=14, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        _add_text_box(slide, left + Inches(0.1), Inches(2.0), Inches(2.0), Inches(0.85),
                      desc, size=12, color=DK2, align=PP_ALIGN.CENTER)
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.35), Inches(4.1),
                                       Inches(1.5), Inches(0.45))
        badge.fill.solid()
        badge.fill.fore_color.rgb = col
        badge.line.fill.background()
        _add_text_box(slide, left + Inches(0.35), Inches(4.15), Inches(1.5), Inches(0.35),
                      status, size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    _card(slide, Inches(0.5), Inches(5.2), Inches(12.3), Inches(1.5), fill=SOFT_BLUE, line=ACCENT1)
    _add_bullets(slide, Inches(0.8), Inches(5.35), Inches(11.8), Inches(1.3), [
        "Near-term: tighten sketch fidelity, add metric dashboard, polish interview UX.",
        "Mid-term: Module A conversational agent + Module E selective refinement.",
        "Long-term: StyleGAN2-ADA + CUFS sketch translation + practitioner validation study.",
    ], size=13, color=DK2, spacing=4)
    _footer(slide, 17)

    # ========== 18. CONCLUSION ==========
    slide = prs.slides.add_slide(blank)
    _header_bar(slide, "Conclusion", "conclusion")
    _add_bullets(slide, Inches(0.7), Inches(1.5), Inches(7.8), Inches(5.0), [
        "Forensic composites remain vital when CCTV / biometrics are unavailable.",
        "We redesign Gen-AI sketching around real CI practice — not one-shot prompts.",
        "Modular pipeline: interview → profile → controllable face → pencil sketch → refine.",
        "50% MVP demonstrates the full path with Streamlit + Stable Diffusion + OpenCV.",
        "Next: StyleGAN control, CUFS sketch translation, evaluation, and user studies.",
        "Positioning: investigator assistive tool with witness-in-the-loop authority.",
    ], size=15, color=DK2, spacing=10)
    ai_img = ASSETS / "img_ai.jpg"
    if ai_img.exists():
        slide.shapes.add_picture(str(ai_img), Inches(8.7), Inches(1.5), Inches(4.2), Inches(5.0))
    _footer(slide, 18)

    # ========== 19. REFERENCES ==========
    slide = prs.slides.add_slide(blank)
    _header_bar(slide, "References", "refs")
    refs = [
        "[1] Soppari et al. (2026). Forensic sketch generation using Gen-AI. WJARR.",
        "[2] Nasir et al. (2019). Text2FaceGAN. IEEE BigMM.",
        "[3] Chen et al. (2019). FTGAN: Fully-trained GAN for text-to-face. arXiv:1904.05729.",
        "[4] Xia et al. (2021). TediGAN. CVPR.",
        "[5] Oza et al. (2021). ST²FG — Semantic text-to-face GAN. arXiv:2107.10756.",
        "[6] Devakumar & Sarath (2023). Forensic sketch to real image using DCGAN. Procedia CS.",
        "[7] Jain et al. (2024). CLIP4Sketch. IJCB.",
        "[8] Galea & Farrugia (2017). Forensic face photo-sketch recognition. IEEE SPL.",
        "[9] Karras et al. (2020). Analyzing and improving StyleGAN. CVPR.",
        "[10] Deng et al. (2019). ArcFace. CVPR.   |   [11] Schroff et al. (2015). FaceNet. CVPR.",
    ]
    _add_bullets(slide, Inches(0.55), Inches(1.3), Inches(12.2), Inches(5.5),
                 refs, size=13, color=DK2, spacing=5)
    _footer(slide, 19)

    # ========== 20. THANK YOU ==========
    slide = prs.slides.add_slide(blank)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = ACCENT5
    bg.line.fill.background()
    band2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(2.6), Inches(13.333), Inches(2.2))
    band2.fill.solid()
    band2.fill.fore_color.rgb = ACCENT1
    band2.line.fill.background()
    _add_text_box(slide, Inches(0.5), Inches(2.9), Inches(12.3), Inches(0.9),
                  "Thank You", size=48, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    _add_text_box(slide, Inches(0.5), Inches(3.85), Inches(12.3), Inches(0.5),
                  "Questions & Discussion", size=22, color=WHITE, align=PP_ALIGN.CENTER)
    _add_text_box(slide, Inches(0.5), Inches(5.2), Inches(12.3), Inches(1.2),
                  "Team: Tusharkanta Behera  ·  Sradha Ram  ·  Sandeep Kumar Swain\n"
                  "Guide: Dr. Sukant Kishoro Bisoy  |  CSE, C.V. Raman Global University\n"
                  "Repo: github.com/Sradha2474/Forensic-Sketch-Generator",
                  size=14, color=WHITE, align=PP_ALIGN.CENTER)

    prs.save(str(OUT))
    print(f"Saved: {OUT}")
    print(f"Slides: {len(prs.slides)}")

if __name__ == "__main__":
    build()
