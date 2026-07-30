from pptx import Presentation
from pathlib import Path
ppt = Path(r"e:\forensic-image-transformation\presentation\Forensic_Sketch_Generator_STAAR_Presentation.pptx")
out = Path(r"e:\forensic-image-transformation\presentation\slide_index.txt")
prs = Presentation(str(ppt))
lines = [f"slides={len(prs.slides)} size_mb={round(ppt.stat().st_size/1024/1024, 2)}"]
for i, slide in enumerate(prs.slides, 1):
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            t = " ".join(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
            if t:
                texts.append(t[:120])
    title = texts[0] if texts else "(no text)"
    lines.append(f"{i:02d}: {title}")
    if len(texts) > 1:
        lines.append(f"    + {texts[1][:100]}")
out.write_text("\n".join(lines), encoding="utf-8")
print("ok", len(prs.slides))
